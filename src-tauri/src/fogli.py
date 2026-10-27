# Halo: operazioni sui fogli Excel (.xlsx) con openpyxl, che mantiene formattazione,
# formule e tutti i fogli. Lo lancia src/fogli.rs così:
#   python3 -I -c <questo script> AZIONE INGRESSO USCITA
# Scrive sempre in USCITA, mai in INGRESSO. Stampa una riga:
#   "OK <riepilogo>"   oppure   "NO <motivo da mostrare>" (ed esce con codice 2).
# Per "cifra" la password arriva sull'ingresso standard, mai nella riga di comando.
import datetime as dt
import json
import re
import sys


def esci(msg, ok=True):
    print(("OK " if ok else "NO ") + msg)
    sys.exit(0 if ok else 2)


def num(n, uno, molti):
    return f"{n} {uno if n == 1 else molti}"


def foglio(wb):
    ws = wb.active
    if ws.max_row < 1 or all(c.value is None for row in ws.iter_rows() for c in row):
        esci(f"Il foglio «{ws.title}» è vuoto.", False)
    return ws


def ha_intestazione(ws):
    """La prima riga è un'intestazione se contiene solo testo."""
    vals = [c.value for c in ws[1] if c.value is not None]
    return bool(vals) and all(isinstance(v, str) for v in vals)


def righe_dati(ws):
    start = 2 if ha_intestazione(ws) else 1
    return list(range(start, ws.max_row + 1))


def vuota(ws, r):
    return all(ws.cell(r, c).value is None or str(ws.cell(r, c).value).strip() == "" for c in range(1, ws.max_column + 1))


def copia_riga(ws, r):
    return [(ws.cell(r, c).value, ws.cell(r, c)._style) for c in range(1, ws.max_column + 1)]


def scrivi_riga(ws, r, dati):
    for c, (v, st) in enumerate(dati, start=1):
        cell = ws.cell(r, c)
        cell.value = v
        cell._style = st


def chiave(v):
    """Ordine naturale: vuoti in fondo, numeri prima del testo, testo senza maiuscole."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return (3, 0)
    if isinstance(v, bool):
        return (0, int(v))
    if isinstance(v, (int, float)):
        return (0, v)
    if isinstance(v, (dt.datetime, dt.date)):
        return (1, v.isoformat())
    return (2, str(v).casefold())


DATE_RE = [
    (re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})$"), "ymd"),
    (re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$"), "dmy"),
]


def come_data(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v
    if not isinstance(v, str):
        return None
    s = v.strip()
    for rx, kind in DATE_RE:
        m = rx.match(s)
        if m:
            a, b, c = (int(x) for x in m.groups())
            y, mo, d = (a, b, c) if kind == "ymd" else (c, b, a)
            try:
                return dt.datetime(y, mo, d)
            except ValueError:
                return None
    return None


def colonna_date(ws, rows):
    for c in range(1, ws.max_column + 1):
        vals = [ws.cell(r, c).value for r in rows if ws.cell(r, c).value is not None]
        if vals and sum(come_data(v) is not None for v in vals) >= max(1, len(vals) * 0.6):
            return c
    return None


def senza_totali(ws, rows):
    """Toglie in fondo le righe vuote e quelle di totale ("Totale", "Total", "Somma"), da non mettere nei grafici."""
    rows = list(rows)
    while rows and (vuota(ws, rows[-1]) or re.match(r"^\s*(totale|totali|total|somma)\b", str(ws.cell(rows[-1], 1).value or ""), re.I)):
        rows.pop()
    return [r for r in rows if not vuota(ws, r)]


def colonne_numeriche(ws, rows):
    out = []
    for c in range(1, ws.max_column + 1):
        vals = [ws.cell(r, c).value for r in rows if ws.cell(r, c).value is not None]
        if vals and sum(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals) >= len(vals) * 0.8:
            out.append(c)
    return out


def main():
    azione, src, dst = sys.argv[1], sys.argv[2], sys.argv[3]

    if azione == "cifra":
        import msoffcrypto
        from msoffcrypto.format.ooxml import OOXMLFile
        pw = sys.stdin.readline().rstrip("\n")
        if not pw:
            esci("Serve una password.", False)
        with open(src, "rb") as f:
            of = OOXMLFile(f)
            if of.is_encrypted():
                esci("Il file ha già una password.", False)
            with open(dst, "wb") as out:
                of.encrypt(pw, out)
        esci("cifrato")

    import openpyxl
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.worksheet.protection import SheetProtection

    try:
        wb = openpyxl.load_workbook(src)
    except Exception as e:  # file cifrato, danneggiato o non .xlsx
        esci(f"Non riesco ad aprire il foglio ({type(e).__name__}). Se ha una password, toglila prima.", False)

    if azione == "json":
        ws = foglio(wb)
        wbv = openpyxl.load_workbook(src, data_only=True)
        wsv = wbv[ws.title]
        rows = list(wsv.iter_rows(values_only=True))
        rows = [r for r in rows if any(v is not None and str(v).strip() for v in r)]
        conv = lambda v: v.isoformat() if isinstance(v, (dt.datetime, dt.date, dt.time)) else v
        if ha_intestazione(ws):
            head = [str(h) if h is not None else f"colonna_{i + 1}" for i, h in enumerate(rows[0])]
            data = [{head[i]: conv(v) for i, v in enumerate(r) if i < len(head)} for r in rows[1:]]
        else:
            data = [[conv(v) for v in r] for r in rows]
        with open(dst, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        esci(num(len(data), "riga", "righe") + f" del foglio «{ws.title}»")

    ws = foglio(wb)
    rows = righe_dati(ws)
    msg = ""

    if azione == "duplicati":
        seen, togli = set(), []
        for r in rows:
            k = tuple(ws.cell(r, c).value for c in range(1, ws.max_column + 1))
            if all(v is None for v in k):
                continue
            if k in seen:
                togli.append(r)
            seen.add(k)
        if not togli:
            esci("Nessuna riga doppia.", False)
        for r in reversed(togli):
            ws.delete_rows(r)
        msg = num(len(togli), "riga doppia tolta", "righe doppie tolte")
    elif azione == "righe-vuote":
        togli = [r for r in range(1, ws.max_row + 1) if vuota(ws, r)]
        # Le righe vuote in fondo non contano: si tolgono solo quelle in mezzo ai dati.
        ultima = max((r for r in range(1, ws.max_row + 1) if not vuota(ws, r)), default=0)
        togli = [r for r in togli if r < ultima]
        if not togli:
            esci("Nessuna riga vuota in mezzo ai dati.", False)
        for r in reversed(togli):
            ws.delete_rows(r)
        msg = num(len(togli), "riga vuota tolta", "righe vuote tolte")
    elif azione == "spazi-in-piu":
        n = 0
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("="):
                    nuovo = " ".join(c.value.split())
                    if nuovo != c.value:
                        c.value = nuovo
                        n += 1
        if not n:
            esci("Nessuno spazio in più da togliere.", False)
        msg = num(n, "cella sistemata", "celle sistemate")
    elif azione == "formato-date":
        n = 0
        for row in ws.iter_rows():
            for c in row:
                d = come_data(c.value)
                if d is None or isinstance(c.value, bool):
                    continue
                ha_ora = isinstance(d, dt.datetime) and (d.hour or d.minute)
                fmt = "DD/MM/YYYY HH:MM" if ha_ora else "DD/MM/YYYY"
                if c.value != d or c.number_format != fmt:
                    c.value = d
                    c.number_format = fmt
                    n += 1
        if not n:
            esci("Nessuna data da sistemare.", False)
        msg = num(n, "data nel formato gg/mm/aaaa", "date nel formato gg/mm/aaaa")
    elif azione in ("a-z", "z-a", "per-data"):
        if len(rows) < 2:
            esci("Ci sono meno di due righe da ordinare.", False)
        col = 1
        if azione == "per-data":
            col = colonna_date(ws, rows)
            if col is None:
                esci("Non trovo una colonna con delle date.", False)
        dati = [copia_riga(ws, r) for r in rows]
        if azione == "per-data":
            dati.sort(key=lambda d: (come_data(d[col - 1][0]) is None, come_data(d[col - 1][0]) or dt.datetime.min))
        else:
            dati.sort(key=lambda d: chiave(d[col - 1][0]), reverse=(azione == "z-a"))
            if azione == "z-a":  # i vuoti restano comunque in fondo
                dati.sort(key=lambda d: chiave(d[col - 1][0])[0] == 3)
        for r, d in zip(rows, dati):
            scrivi_riga(ws, r, d)
        nome = ws.cell(1, col).value if ha_intestazione(ws) else openpyxl.utils.get_column_letter(col)
        msg = num(len(rows), "riga ordinata", "righe ordinate") + f" per «{nome}»"
    elif azione == "filtra":
        ws.auto_filter.ref = ws.dimensions
        msg = f"filtri aggiunti alle intestazioni di «{ws.title}»"
    elif azione in ("barre", "linee", "torta"):
        rows = senza_totali(ws, rows)
        if not ha_intestazione(ws) or len(rows) < 1:
            esci("Serve una prima riga di intestazioni e almeno una riga di dati.", False)
        cols = [c for c in colonne_numeriche(ws, rows) if c != 1]
        if not cols:
            esci("Non trovo colonne di numeri da mettere nel grafico.", False)
        ch = {"barre": BarChart, "linee": LineChart, "torta": PieChart}[azione]()
        if azione == "torta":
            cols = cols[:1]
        # Righe contigue dall'intestazione all'ultima riga di dati (senza il totale finale).
        for c in cols:
            ch.add_data(Reference(ws, min_col=c, min_row=1, max_row=rows[-1]), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=rows[0], max_row=rows[-1]))
        ch.title = ws.cell(1, cols[0]).value if len(cols) == 1 else ws.title
        ch.width, ch.height = 18, 9
        ws.add_chart(ch, f"{openpyxl.utils.get_column_letter(ws.max_column + 2)}2")
        msg = "grafico " + {"barre": "a barre", "linee": "a linee", "torta": "a torta"}[azione] + " accanto ai dati"
    elif azione == "pivot":
        rows = senza_totali(ws, rows)
        if not ha_intestazione(ws):
            esci("Serve una prima riga di intestazioni.", False)
        nums = colonne_numeriche(ws, rows)
        txt = next((c for c in range(1, ws.max_column + 1) if c not in nums), None)
        if txt is None or not nums:
            esci("Serve una colonna di testo (es. categoria) e almeno una di numeri.", False)
        somme = {}
        for r in rows:
            k = ws.cell(r, txt).value
            if k is None:
                continue
            acc = somme.setdefault(k, [0] * len(nums))
            for i, c in enumerate(nums):
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    acc[i] += v
        nome = "Riepilogo"
        while nome in wb.sheetnames:
            nome += " 2"
        out = wb.create_sheet(nome)
        out.append([ws.cell(1, txt).value] + [f"Somma di {ws.cell(1, c).value}" for c in nums])
        for k in sorted(somme, key=chiave):
            out.append([k] + somme[k])
        for cell in out[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        msg = f"foglio «{nome}» con le somme per «{ws.cell(1, txt).value}» ({num(len(somme), 'gruppo', 'gruppi')})"
    elif azione == "blocca-celle":
        for s in wb.worksheets:
            s.protection = SheetProtection(sheet=True, objects=True, scenarios=True)
        msg = num(len(wb.worksheets), "foglio bloccato", "fogli bloccati") + " (senza password)"
    elif azione == "nascondi-formule":
        n = 0
        for s in wb.worksheets:
            for row in s.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        c.protection = openpyxl.styles.Protection(locked=True, hidden=True)
                        n += 1
            s.protection = SheetProtection(sheet=True, objects=True, scenarios=True)
        if not n:
            esci("Il foglio non contiene formule.", False)
        msg = num(n, "formula nascosta", "formule nascoste") + " (fogli protetti senza password)"
    else:
        esci("Opzione non disponibile.", False)

    wb.save(dst)
    esci(msg)


main()
